import io
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional

from docx import Document
from fpdf import FPDF


def format_ms_to_clock(ms: int) -> str:
    """Đổi mili giây sang dạng MM:SS hoặc HH:MM:SS."""
    seconds = ms // 1000
    m, s = divmod(seconds, 60)
    h, m = divmod(m, 60)
    if h > 0:
        return f"{h:02d}:{m:02d}:{s:02d}"
    return f"{m:02d}:{s:02d}"


def format_ms_to_srt_time(ms: int) -> str:
    """Đổi mili giây sang HH:MM:SS,mmm cho SRT."""
    millis = ms % 1000
    seconds = ms // 1000
    m, s = divmod(seconds, 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d},{millis:03d}"


def format_ms_to_vtt_time(ms: int) -> str:
    """Đổi mili giây sang HH:MM:SS.mmm cho VTT."""
    millis = ms % 1000
    seconds = ms // 1000
    m, s = divmod(seconds, 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d}.{millis:03d}"


def resolve_segment_text(seg: Dict[str, Any], text_mode: str) -> str:
    """Chọn lớp văn bản hiển thị theo thứ tự ưu tiên."""
    if text_mode == "raw":
        return seg.get("text_raw", "")
    if text_mode == "cleaned":
        return seg.get("text_cleaned") or seg.get("text_raw", "")
    if text_mode == "edited":
        return seg.get("text_edited") or seg.get("text_cleaned") or seg.get("text_raw", "")
    # "auto"
    return seg.get("text_edited") or seg.get("text_cleaned") or seg.get("text_raw", "")


class ExporterService:
    @staticmethod
    def export_srt(segments: List[Dict[str, Any]], text_mode: str = "auto", include_speakers: bool = True) -> bytes:
        """Xuất phụ đề SRT."""
        lines = []
        for i, seg in enumerate(segments, start=1):
            start = format_ms_to_srt_time(seg["start_ms"])
            end = format_ms_to_srt_time(seg["end_ms"])
            text = resolve_segment_text(seg, text_mode)
            speaker = seg.get("speaker_display") or seg.get("speaker_label")
            prefix = f"{speaker}: " if (include_speakers and speaker) else ""
            lines.append(f"{i}\n{start} --> {end}\n{prefix}{text}\n")
        return "\n".join(lines).encode("utf-8")

    @staticmethod
    def export_vtt(segments: List[Dict[str, Any]], text_mode: str = "auto", include_speakers: bool = True) -> bytes:
        """Xuất phụ đề WebVTT."""
        lines = ["WEBVTT\n"]
        for seg in segments:
            start = format_ms_to_vtt_time(seg["start_ms"])
            end = format_ms_to_vtt_time(seg["end_ms"])
            text = resolve_segment_text(seg, text_mode)
            speaker = seg.get("speaker_display") or seg.get("speaker_label")
            prefix = f"{speaker}: " if (include_speakers and speaker) else ""
            lines.append(f"{start} --> {end}\n{prefix}{text}\n")
        return "\n".join(lines).encode("utf-8")

    @staticmethod
    def export_markdown(
        title: str,
        segments: List[Dict[str, Any]],
        summary: Optional[Dict[str, Any]] = None,
        highlights: Optional[List[Dict[str, Any]]] = None,
        action_items: Optional[List[Dict[str, Any]]] = None,
        text_mode: str = "auto",
        include_speakers: bool = True,
    ) -> bytes:
        """Xuất văn bản Markdown hoàn chỉnh."""
        lines = [f"# {title}\n"]

        if summary:
            lines.append("## Tóm tắt & Biên bản\n")
            if "objective" in summary:
                lines.append(f"**Mục tiêu:** {summary['objective']}\n")
            if "overview" in summary:
                lines.append(f"**Tổng quan:** {summary['overview']}\n")

            if summary.get("decisions"):
                lines.append("### Quyết định đã thông qua\n")
                for d in summary["decisions"]:
                    lines.append(f"- **[{d['timestamp']}]** {d['text']}")
                lines.append("")

            if summary.get("actions"):
                lines.append("### Hành động cần làm (LLM đề xuất)\n")
                for a in summary["actions"]:
                    lines.append(f"- **[{a['timestamp']}]** {a['owner']}: {a['task']} (Hạn: {a.get('deadline', 'N/A')})")
                lines.append("")

        if action_items:
            lines.append("## Việc cần làm (Thủ công)\n")
            for item in action_items:
                ts = f"[{format_ms_to_clock(item['start_ms'])}] " if item.get("start_ms") else ""
                lines.append(f"- {ts}{item.get('owner', '')}: {item['task']} (Hạn: {item.get('deadline', 'N/A')})")
            lines.append("")

        if highlights:
            lines.append("## Đoạn đánh dấu quan trọng\n")
            for hl in highlights:
                ts = f"[{format_ms_to_clock(hl['start_ms'])}]"
                lines.append(f"- {ts} *\"{hl['quote']}\"*")
            lines.append("")

        lines.append("## Nội dung chi tiết\n")
        for seg in segments:
            ts = format_ms_to_clock(seg["start_ms"])
            speaker = seg.get("speaker_display") or seg.get("speaker_label")
            spk_str = f"**{speaker}**: " if (include_speakers and speaker) else ""
            text = resolve_segment_text(seg, text_mode)
            lines.append(f"- `[{ts}]` {spk_str}{text}")

        return "\n".join(lines).encode("utf-8")

    @staticmethod
    def export_docx(
        title: str,
        segments: List[Dict[str, Any]],
        summary: Optional[Dict[str, Any]] = None,
        action_items: Optional[List[Dict[str, Any]]] = None,
        text_mode: str = "auto",
        include_speakers: bool = True,
    ) -> bytes:
        """Xuất tài liệu Word (.docx)."""
        doc = Document()
        doc.add_heading(title, level=0)

        if summary:
            doc.add_heading("Biên bản & Tóm tắt", level=1)
            if "objective" in summary:
                doc.add_paragraph(f"Mục tiêu: {summary['objective']}")
            if "overview" in summary:
                doc.add_paragraph(f"Tổng quan: {summary['overview']}")

        if action_items:
            doc.add_heading("Việc cần làm", level=1)
            for item in action_items:
                ts = f"[{format_ms_to_clock(item['start_ms'])}] " if item.get("start_ms") else ""
                doc.add_paragraph(f"{ts}{item.get('owner', '')}: {item['task']} (Hạn: {item.get('deadline', 'N/A')})", style="List Bullet")

        doc.add_heading("Transcript", level=1)
        for seg in segments:
            ts = format_ms_to_clock(seg["start_ms"])
            speaker = seg.get("speaker_display") or seg.get("speaker_label")
            spk_str = f"{speaker}: " if (include_speakers and speaker) else ""
            text = resolve_segment_text(seg, text_mode)
            doc.add_paragraph(f"[{ts}] {spk_str}{text}")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    @staticmethod
    def export_pdf(
        title: str,
        segments: List[Dict[str, Any]],
        summary: Optional[Dict[str, Any]] = None,
        text_mode: str = "auto",
        include_speakers: bool = True,
    ) -> bytes:
        """Xuất file PDF nhúng font NotoSans để giữ dấu tiếng Việt."""
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)

        font_path = Path(__file__).resolve().parent.parent.parent / "assets" / "fonts" / "NotoSans-Regular.ttf"
        if font_path.exists():
            pdf.add_font("NotoSans", "", str(font_path), uni=True)
            pdf.set_font("NotoSans", size=14)
        else:
            pdf.set_font("Helvetica", size=14)

        pdf.cell(0, 10, title, ln=True, align="C")
        pdf.ln(5)

        if font_path.exists():
            pdf.set_font("NotoSans", size=10)
        else:
            pdf.set_font("Helvetica", size=10)

        for seg in segments:
            ts = format_ms_to_clock(seg["start_ms"])
            speaker = seg.get("speaker_display") or seg.get("speaker_label")
            spk_str = f"{speaker}: " if (include_speakers and speaker) else ""
            text = resolve_segment_text(seg, text_mode)
            line = f"[{ts}] {spk_str}{text}"
            pdf.multi_cell(0, 6, line)
            pdf.ln(1)

        return bytes(pdf.output(dest="S"))
